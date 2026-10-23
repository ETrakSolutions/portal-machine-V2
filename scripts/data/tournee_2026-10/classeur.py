import json,glob,os,re,collections,urllib.parse,sys
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
R=json.load(open('audit/consolide.json',encoding='utf-8'))
TIERS3=('lectura','equipmentworld','forconstructionpros','directindustry','mylittlesalesman','prnewswire','ritchiespecs','constructionequipment.com','kirby-smith')
GLOBAL={'Atlas','CIFA','Everdigm','KCP','Junjin','Mait','Comacchio','Mecalac'}
DOUTE=re.compile(r'douteu|incertain|non confirm|non prouv|global|remorqu|trailer|minier|mining|mine |hors construction|configuration|high reach|straight boom|tunnel|rail|liquide|à confirmer|a confirmer|align|annonc|plus tard|2027|distributeur|concession|dealer',re.I)
rows=[]
for r in R:
    h=urllib.parse.urlparse(r['url']).netloc.lower().replace('www.','')
    txt=' '.join(str(r.get(k) or '') for k in ('note','modele','citation'))
    if r['bd_annees'] and r['statut']=='absent': continue   # faux positif : deja en BD sous ce nom
    if r['statut']=='non_verifiable' or any(t in h for t in TIERS3) or r['page']=='NOM_ABSENT':
        niv='C - Non prouvé'; pq='source tierce, page non vérifiable ou nom introuvable à la relecture'
    elif r['type']=='Excavatrice' and re.match(r'(EX\d{4}|R 99\d|PC\d{4})',str(r['modele'])):
        niv='B - À trancher'; pq='excavatrice minière (> 100 t), périmètre à confirmer'
    elif r['fabricant'] in GLOBAL or not r['fab_en_bd'] or DOUTE.search(txt):
        niv='B - À trancher'; pq=('fabricant absent de ce type en BD' if not r['fab_en_bd'] else 'site global, distribution NA non confirmée' if r['fabricant'] in GLOBAL else 'périmètre ou année à confirmer (voir note)')
    else:
        niv='A - Prêt'; pq=''
    rows.append(dict(r,niveau=niv,pourquoi=pq,domaine=h))
cols=[('niveau','Niveau'),('type','Type'),('fabricant','Fabricant'),('modele','Modèle (nom officiel)'),('statut','Statut'),('annees_a_ajouter','Années à ajouter'),('annee_lancement','Lancement (source)'),('variante_bd_proche','Variante BD proche'),('pourquoi','Pourquoi ce niveau'),('page','Relecture auto de la page'),('domaine','Domaine source'),('url','URL'),('citation','Citation'),('specs_vues','Specs vues'),('note','Note agent'),('decision','Décision Jacquot')]
wb=Workbook(); ws=wb.active; ws.title='Synthèse'
hdr=Font(bold=True,color='FFFFFF'); fill=PatternFill('solid',fgColor='1F3864')
c=collections.Counter((r['type'],r['niveau'][0]) for r in rows); types=sorted({r['type'] for r in rows})
ws.append(['Tournée de vérification des modèles manquants — Portail e-Trak — 2026-10-23']); ws['A1'].font=Font(bold=True,size=13)
ws.append(['Point 1 : Loader complet 2015-2026. Point 2 : autres types, 2025-2026. Rien n\'est importé avant la décision de Jacquot.'])
ws.append([]); ws.append(['Type','A - Prêt','B - À trancher','C - Non prouvé','Total'])
for cell in ws[4]: cell.font=hdr; cell.fill=fill
for t in types: ws.append([t,c[(t,'A')],c[(t,'B')],c[(t,'C')],sum(c[(t,k)] for k in 'ABC')])
ws.append(['TOTAL']+[sum(c[(t,k)] for t in types) for k in 'ABC']+[len(rows)])
for cell in ws[ws.max_row]: cell.font=Font(bold=True)
ws.column_dimensions['A'].width=30
for k,titre in (('A','A - Prêt'),('B','À trancher'),('C','Non prouvé')):
    s=wb.create_sheet(titre); s.append([h for _,h in cols])
    for cell in s[1]: cell.font=hdr; cell.fill=fill
    for r in sorted([x for x in rows if x['niveau'][0]==k],key=lambda x:(x['type'],x['fabricant'],str(x['modele']))):
        s.append([json.dumps(r.get(f),ensure_ascii=False) if isinstance(r.get(f),(list,dict)) else r.get(f) for f,_ in cols])
    for i,w in enumerate([14,22,22,30,11,16,10,18,34,16,24,40,50,30,50,18],1): s.column_dimensions[get_column_letter(i)].width=w
    s.freeze_panes='A2'; s.auto_filter.ref=s.dimensions
s=wb.create_sheet('Anomalies BD'); s.append(['Lot','Fabricant','Modèle en BD','Problème','URL','Décision Jacquot'])
for cell in s[1]: cell.font=hdr; cell.fill=fill
for f in sorted(glob.glob('audit/out_*.json')):
    for a in json.load(open(f,encoding='utf-8')).get('a_verifier',[]):
        s.append([os.path.basename(f)[4:-5],a.get('fabricant'),a.get('modele_bd'),a.get('probleme'),a.get('url'),None])
for i,w in enumerate([14,22,26,90,40,18],1): s.column_dimensions[get_column_letter(i)].width=w
s.freeze_panes='A2'
out=sys.argv[1]; wb.save(out); print('ecrit',out)
for t in types: print(f'{t:28} A={c[(t,"A")]:4} B={c[(t,"B")]:4} C={c[(t,"C")]:4}')
print('TOTAL',len(rows),'A',sum(1 for r in rows if r['niveau'][0]=='A'),'B',sum(1 for r in rows if r['niveau'][0]=='B'),'C',sum(1 for r in rows if r['niveau'][0]=='C'))
